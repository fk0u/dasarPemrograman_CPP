import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCX_PATH = '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/Practical Module 7 Tuple.docx'
doc = docx.Document(DOCX_PATH)

def format_run(run, bold=False, italic=False, size_pt=12, font_name='Times New Roman', color=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_p(cell, text, bold=False, italic=False, size_pt=12, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = align
    run = p.add_run(text)
    format_run(run, bold=bold, italic=italic, size_pt=size_pt)
    return p

def add_img(cell, img_path, width=Inches(5.5), space_after=6):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img_path, width=width)
    return p

# ----------------- TABLE 1: TEST 7.1 -----------------
cell_1 = doc.tables[1].rows[0].cells[0]
# Keep the first paragraph 'Code and Ouput 7.1:'
# Clear any remaining empty paragraphs in cell_1
for p in cell_1.paragraphs[1:]:
    p._element.getparent().remove(p._element)

add_p(cell_1, "A. Source Code (7_1.cpp):", bold=True, space_after=4)
add_img(cell_1, '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/assets/code_7_1.png', width=Inches(5.5))
add_p(cell_1, "B. Program Output:", bold=True, space_after=4)
add_img(cell_1, '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/assets/output_7_1.png', width=Inches(5.2))
add_p(cell_1, "C. Implementation Notes:", bold=True, space_after=2)
add_p(cell_1, "1. Inisialisasi: Tuple dideklarasikan sebagai tuple<string, long long> student(\"Al-Ghani Desta Setyawan\", 2611102441026LL), menyimpan nama mahasiswa bertipe string dan NIM bertipe bilangan bulat 64-bit (long long).")
add_p(cell_1, "2. Akses Elemen: Elemen tuple diakses menggunakan std::get<0>(student) untuk indeks ke-0 (Nama) dan std::get<1>(student) untuk indeks ke-1 (NIM).")
add_p(cell_1, "3. Ukuran Tuple: Fungsi std::tuple_size<decltype(student)>::value mengembalikan ukuran tetap tuple pada waktu kompilasi, yaitu bernilai 2.")

# ----------------- TABLE 3: TEST 7.2 -----------------
cell_3 = doc.tables[3].rows[0].cells[0]
for p in cell_3.paragraphs[1:]:
    p._element.getparent().remove(p._element)

add_p(cell_3, "A. Source Code (7_2.cpp):", bold=True, space_after=4)
add_img(cell_3, '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/assets/code_7_2.png', width=Inches(5.5))
add_p(cell_3, "B. Program Output:", bold=True, space_after=4)
add_img(cell_3, '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/assets/output_7_2.png', width=Inches(5.2))
add_p(cell_3, "C. Implementation Notes:", bold=True, space_after=2)
add_p(cell_3, "1. Packing: Menggunakan std::make_tuple(\"Al-Ghani Desta Setyawan\", 20) untuk mengemas beberapa variabel bernilai beda tipe (nama dan umur) menjadi satu objek tuple secara otomatis.")
add_p(cell_3, "2. Unpacking: Menggunakan std::tie(name, age) = student_data untuk membongkar kembali elemen tuple ke dalam variabel mandiri name dan age yang telah dideklarasikan sebelumnya.")
add_p(cell_3, "3. Nilai yang telah di-unpack dapat langsung digunakan selayaknya variabel biasa.")

# ----------------- TABLE 5: TEST 7.3 -----------------
cell_5 = doc.tables[5].rows[0].cells[0]
for p in cell_5.paragraphs[1:]:
    p._element.getparent().remove(p._element)

add_p(cell_5, "A. Source Code (7_3.cpp):", bold=True, space_after=4)
add_img(cell_5, '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/assets/code_7_3.png', width=Inches(5.5))
add_p(cell_5, "B. Program Output:", bold=True, space_after=4)
add_img(cell_5, '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/assets/output_7_3.png', width=Inches(5.2))
add_p(cell_5, "C. Explanation of Code and Output:", bold=True, space_after=2)
add_p(cell_5, "1. Pembuatan Tuple Terpisah: Program membuat dua tuple yang berbeda spesifikasi:")
add_p(cell_5, "   - Tuple c1 bertipe tuple<const char*, int> berisi data identitas: nama (\"Al-Ghani Desta Setyawan\") dan umur (20).")
add_p(cell_5, "   - Tuple c2 bertipe tuple<int, float> berisi data fisik: tinggi badan (172 cm) dan berat badan (63.5 kg).")
add_p(cell_5, "2. Mekanisme Penggabungan (tuple_cat): Fungsi std::tuple_cat(c1, c2) menggabungkan kedua tuple tersebut secara berurutan menjadi satu tuple komposit baru result bertipe std::tuple<const char*, int, int, float> dengan total 4 elemen.")
add_p(cell_5, "3. Penjelasan Output: Elemen tuple gabungan diakses secara sekuensial dengan get<0>(result) untuk nama, get<1>(result) untuk umur, get<2>(result) untuk tinggi, dan get<3>(result) untuk berat, menghasilkan output profil mahasiswa yang utuh dan terpadu.")

# ----------------- TABLE 7: TEST 7.4 -----------------
cell_7 = doc.tables[7].rows[0].cells[0]
for p in cell_7.paragraphs[1:]:
    p._element.getparent().remove(p._element)

add_p(cell_7, "A. Source Code (7_4.cpp):", bold=True, space_after=4)
add_img(cell_7, '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/assets/code_7_4.png', width=Inches(5.5))
add_p(cell_7, "B. Program Output:", bold=True, space_after=4)
add_img(cell_7, '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/assets/output_7_4.png', width=Inches(5.2))
add_p(cell_7, "C. Explanation of How swap() Works in Tuples:", bold=True, space_after=2)
add_p(cell_7, "1. Kondisi Awal: Dibuat dua tuple person1 (\"Al-Ghani\", 20) dan person2 (\"Desta\", 22) yang memiliki tipe identik tuple<string, int>.")
add_p(cell_7, "2. Mekanisme Kerja swap(): Fungsi person1.swap(person2) melakukan penukaran nilai antar-elemen (element-wise swap) secara berpasangan pada indeks yang bersesuaian, setara dengan menjalankan std::swap(get<0>(person1), get<0>(person2)) dan std::swap(get<1>(person1), get<1>(person2)).")
add_p(cell_7, "3. Syarat Pertukaran: Kedua tuple wajib memiliki tipe dan urutan data yang sama persis serta ukuran yang sama. Jika tipe elemen tidak cocok, compiler akan menolak dengan error static assertion pada waktu kompilasi.")
add_p(cell_7, "4. Hasil Pertukaran: Nilai person1 berubah menjadi (\"Desta\", 22) dan person2 menjadi (\"Al-Ghani\", 20) tanpa perlu membuat objek kontainer sementara yang berat.")

# ----------------- TABLE 8: CONCLUSION -----------------
cell_8 = doc.tables[8].rows[0].cells[0]
# Clear initial empty paragraph
for p in cell_8.paragraphs:
    p._element.getparent().remove(p._element)

add_p(cell_8, "Conclusion of Today's Lesson (Tuple in C++):", bold=True, size_pt=13, space_after=6)
add_p(cell_8, "Berdasarkan praktikum modul 7 yang telah diselesaikan, berikut adalah kesimpulan komprehensif mengenai konsep dan penerapan Tuple dalam bahasa pemrograman C++:")
add_p(cell_8, "1. Definisi dan Konsep Inti:")
add_p(cell_8, "   std::tuple (tersedia sejak C++11 pada header <tuple>) adalah struktur data yang mampu menyimpan kumpulan elemen bertipe data heterogen (berbeda-beda) dengan ukuran tetap yang ditentukan pada waktu kompilasi (compile-time).")
add_p(cell_8, "2. Keunggulan std::tuple dibandingkan std::pair:")
add_p(cell_8, "   Jika std::pair hanya terbatas untuk menampung tepat 2 elemen (.first dan .second), std::tuple dapat menampung berapapun jumlah elemen (N-elemen) tanpa batasan dua nilai, menjadikannya solusi ideal untuk mengembalikan banyak nilai (multiple return values) dari fungsi tanpa perlu mendefinisikan struct khusus.")
add_p(cell_8, "3. Operasi Dasar dan Akses Elemen:")
add_p(cell_8, "   - std::get<index>(tuple): Mengakses nilai elemen pada indeks tertentu (0-indexed). Indeks harus berupa compile-time constant.")
add_p(cell_8, "   - std::tuple_size<decltype(tuple)>::value: Mengetahui banyaknya elemen di dalam tuple secara statis.")
add_p(cell_8, "4. Packing dan Unpacking:")
add_p(cell_8, "   - std::make_tuple(): Membungkus sekumpulan nilai menjadi sebuah objek tuple dengan deduksi tipe otomatis.")
add_p(cell_8, "   - std::tie(): Membongkar kembali elemen tuple ke variabel-variabel penampung terpisah. Kita juga dapat memanfaatkan std::ignore untuk melewatkan elemen tertentu yang tidak diperlukan.")
add_p(cell_8, "5. Penggabungan dan Pertukaran (tuple_cat & swap):")
add_p(cell_8, "   - std::tuple_cat(): Menggabungkan dua atau lebih tuple berbeda tipe menjadi satu tuple terpadu dengan ukuran kumulatif.")
add_p(cell_8, "   - .swap(): Menukar isi dari dua tuple berukuran dan bertipe sama secara efisien dengan semantik pemindahan (move semantics).")
add_p(cell_8, "6. Relevansi pada Standar Modern C++ (C++17):")
add_p(cell_8, "   Pada C++17 ke atas, tuple semakin fleksibel dan ringkas dengan hadirnya fitur Structured Binding (seperti auto [name, age] = student_data;), yang membuat sintaksis penguraian data tuple jauh lebih ekspresif dan mudah dipelihara.")

doc.save(DOCX_PATH)
print("Practical Module 7 Tuple.docx populated and saved successfully!")
