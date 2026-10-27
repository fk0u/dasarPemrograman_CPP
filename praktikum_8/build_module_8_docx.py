import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC_DOCX = '/Users/kou/Project/dasarPemrograman_CPP/praktikum_7/Practical Module 7 Tuple.docx'
DST_DOCX = '/Users/kou/Project/dasarPemrograman_CPP/praktikum_8/Practical Module 8 Set and Map.docx'
ASSETS_DIR = '/Users/kou/Project/dasarPemrograman_CPP/praktikum_8/assets'

shutil.copyfile(SRC_DOCX, DST_DOCX)
doc = docx.Document(DST_DOCX)

def format_run(run, bold=False, italic=False, size_pt=12, font_name='Times New Roman', color=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_p(cell, text, bold=False, italic=False, size_pt=12, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = align
    run = p.add_run(text)
    format_run(run, bold=bold, italic=italic, size_pt=size_pt)
    return p

def add_img(cell, img_path, width=Inches(5.5), space_after=6):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(img_path, width=width)
    return p

# 1. Update Paragraphs (Cover and Headers)
for p in doc.paragraphs:
    if "Practicum Report 7 : Tuple" in p.text:
        p.text = p.text.replace("Practicum Report 7 : Tuple", "Practicum Report 8 : Set & Map")
        if p.runs: format_run(p.runs[0], bold=True, size_pt=14)
    elif "Practicum Report 7:" in p.text:
        p.text = p.text.replace("Practicum Report 7:", "Practicum Report 8:")
        if p.runs: format_run(p.runs[0], bold=True, size_pt=14)
    elif p.text.strip() == "Tuple":
        p.text = "Set & Map"
        if p.runs: format_run(p.runs[0], bold=True, size_pt=18)
    elif "Implementing tuples in C++" in p.text:
        p.text = p.text.replace("Implementing tuples in C++", "Implementation the use of set & map in C++.")
        if p.runs: format_run(p.runs[0], size_pt=12)
    elif "Basic Operation on Tuples" in p.text:
        p.text = "1. Basic Operation on Set"
        if p.runs: format_run(p.runs[0], bold=True, size_pt=12)
    elif "Tuple Packing and Unpacking" in p.text:
        p.text = "2. Advance Operation on Set"
        if p.runs: format_run(p.runs[0], bold=True, size_pt=12)
    elif "Concatenating Tuples" in p.text:
        p.text = "3. Basic Operation on Map"
        if p.runs: format_run(p.runs[0], bold=True, size_pt=12)
    elif "Swapping Tuples" in p.text:
        p.text = "4. Using Functions on Map"
        if p.runs: format_run(p.runs[0], bold=True, size_pt=12)

# ----------------- TABLE 0: TEST 8.1 INSTRUCTIONS -----------------
cell_0 = doc.tables[0].rows[0].cells[0]
for p in list(cell_0.paragraphs):
    p._element.getparent().remove(p._element)

add_p(cell_0, "Test 8.1: Basic Operation on Set", bold=True, size_pt=12, space_after=4)
add_p(cell_0, "Instructions:", bold=True, size_pt=12, space_after=2)
add_p(cell_0, "1. Create a Set containing fruit types.")
add_p(cell_0, "2. Add a new fruit using insert().")
add_p(cell_0, "3. Remove a fruit using erase().")
add_p(cell_0, "4. Display the contents of the Set after each operation.")

# ----------------- TABLE 1: TEST 8.1 CODE & OUTPUT -----------------
cell_1 = doc.tables[1].rows[0].cells[0]
for p in list(cell_1.paragraphs):
    p._element.getparent().remove(p._element)

add_p(cell_1, "Code and Output 8.1:", bold=True, size_pt=12, space_after=4)
add_p(cell_1, "A. Source Code (8_1.cpp):", bold=True, space_after=4)
add_img(cell_1, os.path.join(ASSETS_DIR, 'code_8_1.png'), width=Inches(5.5))
add_p(cell_1, "B. Program Output:", bold=True, space_after=4)
add_img(cell_1, os.path.join(ASSETS_DIR, 'output_8_1.png'), width=Inches(5.2))
add_p(cell_1, "C. Implementation Notes:", bold=True, space_after=2)
add_p(cell_1, "1. Inisialisasi: std::set<string> fruits otomatis mengurutkan elemen string secara ascending (alfabetis) dan menjamin keunikan nilai.")
add_p(cell_1, "2. Operasi insert(): fruits.insert(\"Mango\") menyisipkan elemen baru, yang langsung ditempatkan pada posisi terurut yang sesuai.")
add_p(cell_1, "3. Operasi erase(): fruits.erase(\"Banana\") menghapus elemen spesifik secara efisien berdasarkan nilai kuncinya.")

# ----------------- TABLE 2: TEST 8.2 INSTRUCTIONS -----------------
cell_2 = doc.tables[2].rows[0].cells[0]
for p in list(cell_2.paragraphs):
    p._element.getparent().remove(p._element)

add_p(cell_2, "Test 8.2: Advance Operation on Set", bold=True, size_pt=12, space_after=4)
add_p(cell_2, "Instructions:", bold=True, size_pt=12, space_after=2)
add_p(cell_2, "1. Create two Sets containing numbers.")
add_p(cell_2, "2. Display the size of each Set.")
add_p(cell_2, "3. Remove all elements from one of the Sets using clear().")
add_p(cell_2, "4. Check whether the Set is empty using empty().")

# ----------------- TABLE 3: TEST 8.2 CODE & OUTPUT -----------------
cell_3 = doc.tables[3].rows[0].cells[0]
for p in list(cell_3.paragraphs):
    p._element.getparent().remove(p._element)

add_p(cell_3, "Code and Output 8.2:", bold=True, size_pt=12, space_after=4)
add_p(cell_3, "A. Source Code (8_2.cpp):", bold=True, space_after=4)
add_img(cell_3, os.path.join(ASSETS_DIR, 'code_8_2.png'), width=Inches(5.5))
add_p(cell_3, "B. Program Output:", bold=True, space_after=4)
add_img(cell_3, os.path.join(ASSETS_DIR, 'output_8_2.png'), width=Inches(5.2))
add_p(cell_3, "C. Implementation Notes:", bold=True, space_after=2)
add_p(cell_3, "1. Ukuran Set: Fungsi size() mengembalikan jumlah elemen yang tersimpan secara dinamis (Set1 berukuran 4, Set2 berukuran 5).")
add_p(cell_3, "2. Pengosongan Set: set1.clear() menghapus seluruh elemen di dalam set1 sehingga ukurannya menjadi 0.")
add_p(cell_3, "3. Pemeriksaan Kekosongan: Fungsi empty() mengembalikan nilai boolean true jika set kosong dan false jika masih memiliki elemen.")

# ----------------- TABLE 4: TEST 8.3 INSTRUCTIONS -----------------
cell_4 = doc.tables[4].rows[0].cells[0]
for p in list(cell_4.paragraphs):
    p._element.getparent().remove(p._element)

add_p(cell_4, "Test 8.3: Basic Operation on Map", bold=True, size_pt=12, space_after=4)
add_p(cell_4, "Instructions:", bold=True, size_pt=12, space_after=2)
add_p(cell_4, "1. Create a Map containing student information (name as key, age as value).")
add_p(cell_4, "2. Add new data using insert().")
add_p(cell_4, "3. Delete specific data using erase().")
add_p(cell_4, "4. Retrieve values based on specific keys using at().")

# ----------------- TABLE 5: TEST 8.3 CODE & OUTPUT -----------------
cell_5 = doc.tables[5].rows[0].cells[0]
for p in list(cell_5.paragraphs):
    p._element.getparent().remove(p._element)

add_p(cell_5, "Code and Output 8.3:", bold=True, size_pt=12, space_after=4)
add_p(cell_5, "A. Source Code (8_3.cpp):", bold=True, space_after=4)
add_img(cell_5, os.path.join(ASSETS_DIR, 'code_8_3.png'), width=Inches(5.5))
add_p(cell_5, "B. Program Output:", bold=True, space_after=4)
add_img(cell_5, os.path.join(ASSETS_DIR, 'output_8_3.png'), width=Inches(5.2))
add_p(cell_5, "C. Implementation Notes:", bold=True, space_after=2)
add_p(cell_5, "1. Struktur Pasangan Kunci-Nilai: std::map<string, int> menyimpan pasangan key unik (nama) yang dipetakan ke nilainya (umur).")
add_p(cell_5, "2. Operasi insert(): Menambahkan pasangan baru {\"Kara\", 19} dengan aman.")
add_p(cell_5, "3. Operasi erase(): Menghapus pasangan data dengan kunci \"Baskara\".")
add_p(cell_5, "4. Akses dengan at(): students.at(\"Al-Ghani\") mengambil nilai umur mahasiswa dengan pengecekan batas (bounds checking).")

# ----------------- TABLE 6: TEST 8.4 INSTRUCTIONS -----------------
cell_6 = doc.tables[6].rows[0].cells[0]
for p in list(cell_6.paragraphs):
    p._element.getparent().remove(p._element)

add_p(cell_6, "Test 8.4: Using Functions on Map", bold=True, size_pt=12, space_after=4)
add_p(cell_6, "Instructions:", bold=True, size_pt=12, space_after=2)
add_p(cell_6, "1. Create a Map containing several students with their names as keys and student ages as values.")
add_p(cell_6, "2. Display all keys, values, and key-value pairs using first and second.")
add_p(cell_6, "3. Check whether a specific key exists in the Map using count().")
add_p(cell_6, "4. Delete all student data using clear().")
add_p(cell_6, "5. Display the size of the Map using size().")
add_p(cell_6, "6. Check whether the Map is empty using empty().")

# ----------------- TABLE 7: TEST 8.4 CODE & OUTPUT -----------------
cell_7 = doc.tables[7].rows[0].cells[0]
for p in list(cell_7.paragraphs):
    p._element.getparent().remove(p._element)

add_p(cell_7, "Code and Output 8.4:", bold=True, size_pt=12, space_after=4)
add_p(cell_7, "A. Source Code (8_4.cpp):", bold=True, space_after=4)
add_img(cell_7, os.path.join(ASSETS_DIR, 'code_8_4.png'), width=Inches(5.5))
add_p(cell_7, "B. Program Output:", bold=True, space_after=4)
add_img(cell_7, os.path.join(ASSETS_DIR, 'output_8_4.png'), width=Inches(5.2))
add_p(cell_7, "C. Implementation Notes:", bold=True, space_after=2)
add_p(cell_7, "1. Iterasi first & second: Setiap elemen map berupa std::pair; .first merepresentasikan kunci (nama) dan .second merepresentasikan nilai (umur).")
add_p(cell_7, "2. Validasi Kunci dengan count(): students.count(\"Desta\") bernilai 1 jika kunci ada, dan 0 jika tidak ada.")
add_p(cell_7, "3. Pengosongan & Status: clear() mengosongkan seluruh isi map, divalidasi dengan size() == 0 dan empty() == true.")

# ----------------- TABLE 8: CONCLUSION -----------------
cell_8 = doc.tables[8].rows[0].cells[0]
for p in list(cell_8.paragraphs):
    p._element.getparent().remove(p._element)

add_p(cell_8, "Conclusion of Today's Lesson:", bold=True, size_pt=13, space_after=6)
add_p(cell_8, "From today’s practicum, I learned how to implement and manipulate associative containers (std::set and std::map) in C++:")
add_p(cell_8, "1. Understanding std::set: A set is an associative container that stores unique elements in automatically sorted order. I learned to insert new elements with insert(), remove elements with erase(), inspect the collection size with size(), and check empty status with empty().")
add_p(cell_8, "2. Understanding std::map: A map stores elements as key-value pairs (std::pair) where each key is unique and sorted. I learned to iterate through elements accessing keys via .first and values via .second.")
add_p(cell_8, "3. Key Operations in Map: I practiced safe value retrieval with at() (which provides bounds checking), checking for key existence cleanly using count(), and clearing all data using clear().")
add_p(cell_8, "Overall, std::set and std::map are powerful standard library containers that make searching, sorting, and data association efficient and clean in C++ programming.")

doc.save(DST_DOCX)
print("Practical Module 8 Set and Map.docx created and saved successfully!")
